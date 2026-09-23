"""Date findings stay on their own tab, including old Main/Date dual routes."""
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.styles import Border, PatternFill, Side

from generate_reports.create_trackers import CreateTrackers


NETWORKS = ['PRESCIENT', 'PRONET']


def _mapping():
    return {
        'subject': 'Participant', 'cohort': 'Cohort',
        'displayed_timepoint': 'Timepoint', 'displayed_form': 'Form',
        'flag_count': 'Flag Count', 'error_message': 'Flags',
        'var_translations': 'Translations',
        'time_since_last_detection': 'Days Since Detected',
        'date_resolved': 'Date Resolved', 'manually_resolved': 'Manually Resolved',
        'comments': 'Network Comments', 'site_comments': 'Site Comments',
        'priority_item': 'Priority Item',
    }


def _row(network, reports, message):
    return {
        'network': network, 'subject': 'PI001' if network == 'PRESCIENT' else 'KC001',
        'cohort': 'CHR', 'reports': reports, 'displayed_timepoint': 'month1',
        'displayed_form': 'visit_form', 'affected_forms': 'visit_form',
        'displayed_variable': 'visit_date', 'affected_variables': 'visit_date',
        'error_message': message, 'currently_resolved': False,
        'manually_resolved': '', 'dates_resolved': '', 'var_translations': '',
        'time_since_last_detection': 0, 'priority_item': False,
        'comments': 'network note', 'site_comments': 'site note',
    }


def _date_message(days):
    return (f'visit_date : Visit date at month1 (2025-01-01) is {days} '
            'day(s) before visit_date (2025-01-11) at the earlier '
            'timepoint baseline.')


def _generator(tmp_path, network):
    tracker = object.__new__(CreateTrackers)
    site = 'PI' if network == 'PRESCIENT' else 'KC'
    tracker.utils = SimpleNamespace(
        pipeline_networks=(network,),
        site_full_name_translations={'PI': 'Pittsburgh', 'KC': 'Kansas City'},
        can_be_float=lambda value: str(value).replace('.', '', 1).isdigit())
    tracker.all_sites = {network: [site]}
    tracker.all_reports = ['Main Report', 'Date Report']
    tracker.site_reports = ['Main Report', 'Date Report']
    tracker.all_report_df = {}
    tracker.dropbox_output_path = str(tmp_path) + '/'
    tracker.melbourne_ras = {}
    tracker.formatted_column_names = {
        network: {'combined': _mapping(), 'sites': _mapping()}}
    tracker.colors = {
        name: PatternFill(fill_type='solid', fgColor='FFFFFF')
        for name in ('green', 'yellow', 'blue', 'orange', 'red', 'grey', 'pink')}
    tracker.thin_border = Border(left=Side(style='thin'))
    return tracker


def test_main_selection_excludes_exact_date_routes_without_mutating_history():
    frame = pd.DataFrame({
        'reports': ['Date Report', 'Main Report | Date Report',
                    'Date Report | Main Report', 'Main Report',
                    'Main Report | Old Date Report',
                    'Main Report | Date Report Extended',
                    'Main Report Extended', None],
        'error_message': list('abcdefgh'),
    }, index=[10, 20, 30, 40, 50, 60, 70, 80])
    original = frame.copy(deep=True)

    selected = CreateTrackers._select_main_report_rows(frame)

    assert selected.index.tolist() == [40, 50, 60]
    assert selected['error_message'].tolist() == ['d', 'e', 'f']
    pd.testing.assert_frame_equal(frame, original)
    selected.loc[40, 'reports'] = 'presentation edit'
    pd.testing.assert_frame_equal(frame, original)


@pytest.mark.parametrize('network', NETWORKS)
@pytest.mark.parametrize('ordinary_main', [True, False])
def test_regeneration_removes_date_findings_from_combined_and_site_main(
        tmp_path, network, ordinary_main):
    tracker = _generator(tmp_path, network)
    site_code, site_name = (
        ('PI', 'Pittsburgh') if network == 'PRESCIENT'
        else ('KC', 'Kansas City'))
    paths = [
        tmp_path / network / 'combined' / f'{network}_Output_V2.xlsx',
        tmp_path / network / site_name / f'{network}_{site_code}_Output_V2.xlsx',
    ]
    # Simulate a workbook produced before Date/Main separation, including a
    # Main sheet that becomes empty when the last ordinary Main flag resolves.
    for path in paths:
        path.parent.mkdir(parents=True, exist_ok=True)
        with pd.ExcelWriter(path, engine='openpyxl') as writer:
            pd.DataFrame({'Participant': ['old'], 'Flags': ['old date flag']}).to_excel(
                writer, sheet_name='Main Report', index=False)
            pd.DataFrame({'Participant': ['old'], 'Flags': ['old date flag']}).to_excel(
                writer, sheet_name='Date Report', index=False)
            pd.DataFrame({'Note': ['keep this operator note']}).to_excel(
                writer, sheet_name='Operator Notes', index=False)

    date_messages = [_date_message(10), _date_message(20)]
    rows = [
        _row(network, 'Date Report', date_messages[0]),
        _row(network, 'Main Report | Date Report', date_messages[1]),
    ]
    ordinary_messages = ['field_a : Variable is blank.', 'field_b : Out of range.']
    if ordinary_main:
        rows.extend(_row(network, 'Main Report', message)
                    for message in ordinary_messages)
    # All rows deliberately share subject/timepoint/form. Filtering after the
    # real grouping stage would incorrectly retain their Date messages in Main.
    tracker.combined_tracker = pd.DataFrame(rows)
    original = tracker.combined_tracker.copy(deep=True)
    tracker.collect_new_reports()
    tracker.generate_reports()
    tracker.blank_stale_sheets()
    pd.testing.assert_frame_equal(tracker.combined_tracker, original)

    for path in paths:
        workbook = load_workbook(path)
        try:
            main_rows = list(workbook['Main Report'].values)
            assert len(main_rows) == (2 if ordinary_main else 1)
            if ordinary_main:
                headers, values = main_rows
                assert values[headers.index('Flag Count')] == 2
                flags = values[headers.index('Flags')]
                assert all(message in flags for message in ordinary_messages)
                assert all(message not in flags for message in date_messages)
            date_rows = list(workbook['Date Report'].values)
            assert len(date_rows) == 2
            headers, values = date_rows
            assert values[headers.index('Flag Count')] == 2
            flags = values[headers.index('Flags')]
            assert all(message in flags for message in date_messages)
            assert all(message not in flags for message in ordinary_messages)
            assert workbook['Operator Notes']['A2'].value == 'keep this operator note'
        finally:
            workbook.close()
