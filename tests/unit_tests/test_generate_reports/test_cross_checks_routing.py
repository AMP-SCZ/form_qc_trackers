"""Regression coverage for the dedicated Cross Checks tracker tab."""

from io import BytesIO
from types import SimpleNamespace

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Border, PatternFill, Side

from generate_reports.calculate_resolved_errors import CalculateResolvedErrors
from generate_reports.create_trackers import CreateTrackers


def _column_mapping():
    return {
        "subject": "Participant",
        "cohort": "Cohort",
        "displayed_timepoint": "Timepoint",
        "displayed_form": "Form",
        "flag_count": "Flag Count",
        "error_message": "Flags",
        "var_translations": "Translations",
        "time_since_last_detection": "Days Since Detected",
        "date_resolved": "Date Resolved",
        "manually_resolved": "Manually Resolved",
        "comments": "Network Comments",
        "site_comments": "Site Comments",
        "priority_item": "Priority Item",
    }


def _workbook_bytes(sheets):
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for sheet_name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=sheet_name, index=False)
    return buffer.getvalue()










def test_cross_checks_manual_resolution_and_comments_round_trip():
    calculator = object.__new__(CalculateResolvedErrors)
    calculator.utils = SimpleNamespace(pipeline_networks=("PRONET", "PRESCIENT"), 
        reverse_dictionary=lambda mapping: {
            value: key for key, value in mapping.items()})
    calculator._corrupt_workbooks = []

    current = pd.DataFrame([{
        "network": "PRONET",
        "subject": "KC001",
        "displayed_timepoint": "baseline",
        "displayed_form": "current_health_status",
        "error_message": "[CROSS-QC-005] Cross-form review: example",
        "reports": "Cross Checks",
        "manually_resolved": "",
        "comments": "",
    }])
    sheet = pd.DataFrame([{
        "Participant": "KC001",
        "Timepoint": "baseline",
        "Form": "current_health_status",
        "Flags": "[CROSS-QC-005] Cross-form review: example",
        "Manually Resolved": "yes",
        "Network Comments": "Reviewed against source record",
    }])

    result = calculator.read_dropbox_data(
        current, _column_mapping(), ["manually_resolved", "comments"],
        "unused/PRONET_Output.xlsx", None, "PRONET",
        ["Cross Checks"],
        preloaded_data=_workbook_bytes({"Cross Checks": sheet}),
    ).iloc[0]

    assert result["manually_resolved"] == "yes"
    assert result["comments"] == "Reviewed against source record"




def test_melbourne_ra_writer_and_reader_use_identical_sanitized_path():
    ra_name = "Dr Jane Doe/RA #1"
    tracker_root = "/trackers/"

    writer = object.__new__(CreateTrackers)
    writer.dropbox_output_path = tracker_root
    writer.melbourne_ras = {ra_name: ["ME001"]}
    writer_writes = []
    writer.format_excl_sheet = (
        lambda frame, report, folder, filename:
        writer_writes.append((frame.copy(), report, folder, filename)))
    writer.loop_ras(
        "PRESCIENT", "Melbourne", "Cross Checks",
        pd.DataFrame({"Participant": ["ME001", "ME999"]}),
    )

    assert len(writer_writes) == 1
    written_frame, _report, folder, filename = writer_writes[0]
    writer_path = folder + filename
    assert written_frame["Participant"].tolist() == ["ME001"]
    assert writer_path == (
        "/trackers/PRESCIENT/Melbourne/Dr_Jane_Doe_RA__1/"
        "PRESCIENT_Melbourne_Output_V2.xlsx")

    class FakeDropbox:
        def files_list_folder(self, _path):
            return SimpleNamespace(entries=[SimpleNamespace(name="PRESCIENT")])

        def files_download(self, _path):
            return None, SimpleNamespace(content=b"workbook")

    reader = object.__new__(CalculateResolvedErrors)
    reader.utils = SimpleNamespace(pipeline_networks=("PRONET", "PRESCIENT"), 
        collect_dropbox_credentials=lambda: FakeDropbox(),
        all_sites={"PRESCIENT": ["ME"]},
        site_full_name_translations={"ME": "Melbourne"},
    )
    reader.dropbox_path = tracker_root
    reader.formatted_column_names = {
        "PRESCIENT": {
            "combined": _column_mapping(),
            "sites": _column_mapping(),
        }}
    reader.melbourne_ras = {ra_name: ["ME001"]}
    reader._read_current = lambda: pd.DataFrame({"subject": ["ME001"]})
    written_current = []
    reader._write_current = lambda frame: written_current.append(frame.copy())
    reader.check_dbx_file_exists = lambda dbx, path: True
    reader_calls = []

    def capture_read(frame, col_names, columns_to_read, path, dbx, network,
                     reports_to_read, **kwargs):
        reader_calls.append((path, list(columns_to_read)))
        if path.endswith("/PRESCIENT_ME_Output_V2.xlsx"):
            return frame.assign(site_comments="broad site comment")
        if path == writer_path:
            return frame.assign(site_comments="RA-specific comment")
        return frame

    reader.read_dropbox_data = capture_read
    reader.loop_dropbox_files()

    melbourne_read_paths = [
        path for path, _columns in reader_calls if "/Melbourne/" in path]
    assert melbourne_read_paths == [
        "/trackers/PRESCIENT/Melbourne/PRESCIENT_ME_Output_V2.xlsx",
        writer_path,
    ]
    assert all(columns == ["site_comments"]
               for path, columns in reader_calls if "/Melbourne/" in path)
    # Broad site data is a fallback; the more-specific RA surface wins.
    assert written_current[0].iloc[0]["site_comments"] == "RA-specific comment"


def test_cross_check_id_round_trip_survives_display_text_changes():
    calculator = object.__new__(CalculateResolvedErrors)
    calculator.utils = SimpleNamespace(pipeline_networks=("PRONET", "PRESCIENT"), 
        reverse_dictionary=lambda mapping: {
            value: key for key, value in mapping.items()})
    calculator._corrupt_workbooks = []

    current = pd.DataFrame([{
        "network": "PRONET",
        "subject": "KC001",
        "displayed_timepoint": "baseline",
        "displayed_form": "old_display_form",
        "error_message": "[CROSS-QC-005] Cross-form review: old wording.",
        "check_id": "CROSS-QC-005",
        "reports": "Cross Checks",
        "manually_resolved": "",
        "comments": "",
    }])
    # The form and human-readable message changed, but Check ID did not.
    sheet = pd.DataFrame([{
        "Participant": "KC001",
        "Timepoint": "baseline",
        "Form": "new_display_form",
        "Flags": "[CROSS-QC-005] Cross-form review: clearer wording.",
        "Check ID": "CROSS-QC-005",
        "Manually Resolved": "yes",
        "Network Comments": "Reviewed after wording update",
    }])

    result = calculator.read_dropbox_data(
        current, _column_mapping(), ["manually_resolved", "comments"],
        "unused/PRONET_Output.xlsx", None, "PRONET",
        ["Cross Checks"],
        preloaded_data=_workbook_bytes({"Cross Checks": sheet}),
    ).iloc[0]

    assert result["manually_resolved"] == "yes"
    assert result["comments"] == "Reviewed after wording update"
    assert not any(str(column).startswith("__cross_")
                   for column in result.index)


def test_explicit_clear_token_clears_but_blank_cells_are_noops():
    calculator = object.__new__(CalculateResolvedErrors)
    calculator.utils = SimpleNamespace(pipeline_networks=("PRONET", "PRESCIENT"), 
        reverse_dictionary=lambda mapping: {
            value: key for key, value in mapping.items()})
    calculator._corrupt_workbooks = []

    current = pd.DataFrame([{
        "network": "PRONET",
        "subject": "KC001",
        "displayed_timepoint": "baseline",
        "displayed_form": "current_health_status",
        "error_message": "[CROSS-QC-005] Cross-form review: example",
        "check_id": "CROSS-QC-005",
        "reports": "Cross Checks",
        "manually_resolved": "yes",
        "comments": "keep this comment",
    }])
    identity = {
        "Participant": "KC001",
        "Timepoint": "baseline",
        "Form": "current_health_status",
        "Flags": "[CROSS-QC-005] Cross-form review: example",
        "Check ID": "CROSS-QC-005",
    }
    blank_sheet = pd.DataFrame([{
        **identity, "Manually Resolved": "   ", "Network Comments": "\t"}])
    blank_result = calculator.read_dropbox_data(
        current, _column_mapping(), ["manually_resolved", "comments"],
        "unused.xlsx", None, "PRONET", ["Cross Checks"],
        preloaded_data=_workbook_bytes({"Cross Checks": blank_sheet}),
    ).iloc[0]
    assert blank_result["manually_resolved"] == "yes"
    assert blank_result["comments"] == "keep this comment"

    clear_sheet = pd.DataFrame([{
        **identity,
        "Manually Resolved": " [clear] ",
        "Network Comments": "[CLEAR]",
    }])
    cleared = calculator.read_dropbox_data(
        current, _column_mapping(), ["manually_resolved", "comments"],
        "unused.xlsx", None, "PRONET", ["Cross Checks"],
        preloaded_data=_workbook_bytes({"Cross Checks": clear_sheet}),
    ).iloc[0]
    assert cleared["manually_resolved"] == ""
    assert cleared["comments"] == ""


def test_cross_check_workbook_cannot_update_the_other_network():
    calculator = object.__new__(CalculateResolvedErrors)
    calculator.utils = SimpleNamespace(pipeline_networks=("PRONET", "PRESCIENT"), 
        reverse_dictionary=lambda mapping: {
            value: key for key, value in mapping.items()})
    calculator._corrupt_workbooks = []

    shared = {
        "subject": "SHARED001",
        "displayed_timepoint": "baseline",
        "displayed_form": "form_a",
        "error_message": "[CROSS-QC-005] Cross-form review: example",
        "check_id": "CROSS-QC-005",
        "reports": "Cross Checks",
        "manually_resolved": "",
        "comments": "",
    }
    current = pd.DataFrame([
        {**shared, "network": "PRONET"},
        {**shared, "network": "PRESCIENT"},
    ])
    sheet = pd.DataFrame([{
        "Participant": "SHARED001",
        "Timepoint": "baseline",
        "Form": "form_a",
        "Flags": "[CROSS-QC-005] Cross-form review: example",
        "Check ID": "CROSS-QC-005",
        "Manually Resolved": "yes",
        "Network Comments": "PRONET reviewer",
    }])

    result = calculator.read_dropbox_data(
        current, _column_mapping(), ["manually_resolved", "comments"],
        "unused/PRONET_Output.xlsx", None, "PRONET", ["Cross Checks"],
        preloaded_data=_workbook_bytes({"Cross Checks": sheet}),
    ).set_index("network")

    assert result.loc["PRONET", "manually_resolved"] == "yes"
    assert result.loc["PRONET", "comments"] == "PRONET reviewer"
    assert result.loc["PRESCIENT", "manually_resolved"] == ""
    assert result.loc["PRESCIENT", "comments"] == ""


def test_conflicting_check_id_and_message_prefix_fail_closed():
    calculator = object.__new__(CalculateResolvedErrors)
    calculator.utils = SimpleNamespace(pipeline_networks=("PRONET", "PRESCIENT"), 
        reverse_dictionary=lambda mapping: {
            value: key for key, value in mapping.items()})
    calculator._corrupt_workbooks = []
    current = pd.DataFrame([{
        "network": "PRONET",
        "subject": "KC001",
        "displayed_timepoint": "baseline",
        "displayed_form": "form_a",
        "error_message": "[CROSS-QC-005] Cross-form review: example",
        "check_id": "CROSS-QC-005",
        "reports": "Cross Checks",
        "comments": "original",
    }])
    sheet = pd.DataFrame([{
        "Participant": "KC001",
        "Timepoint": "baseline",
        "Form": "form_a",
        "Flags": "[CROSS-QC-005] Cross-form review: example",
        # A hand-edited machine ID disagrees with the message prefix. Applying
        # it by either identity could target the wrong rule, so ignore it.
        "Check ID": "CROSS-QC-006",
        "Network Comments": "must not apply",
    }])

    result = calculator.read_dropbox_data(
        current, _column_mapping(), ["comments"],
        "unused.xlsx", None, "PRONET", ["Cross Checks"],
        preloaded_data=_workbook_bytes({"Cross Checks": sheet}),
    )

    assert len(result) == 1
    assert result.iloc[0]["comments"] == "original"


def test_cross_checks_workbook_writes_and_hides_stable_check_ids(tmp_path):
    trackers = object.__new__(CreateTrackers)
    trackers.formatted_column_names = {
        "PRONET": {"combined": _column_mapping()}}

    raw = pd.DataFrame([{
        "subject": "KC001",
        "cohort": "CHR",
        "displayed_timepoint": "baseline",
        "displayed_form": "form_a",
        "currently_resolved": False,
        "manually_resolved": "",
        "dates_resolved": "",
        "var_translations": "translation",
        "error_message": "[CROSS-QC-005] First",
        "time_since_last_detection": 2,
        "priority_item": False,
        "comments": "",
        "site_comments": "",
        "reports": "Cross Checks",
        # Legacy resolved rows can have the message prefix but a blank column;
        # the writer must backfill the hidden machine identity.
        "check_id": "",
    }])
    formatted = trackers.convert_to_shared_format(raw, "PRONET")
    assert formatted["Check ID"].tolist() == ["CROSS-QC-005"]

    # Reuse a fully initialized styling fixture shape for the actual workbook.
    trackers.utils = SimpleNamespace(pipeline_networks=("PRONET", "PRESCIENT"), can_be_float=lambda value: False)
    trackers.colors = {
        name: PatternFill(
            start_color="ededed", end_color="ededed", fill_type="solid")
        for name in ["green", "yellow", "blue", "orange", "red", "grey", "pink"]
    }
    trackers.thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"))
    trackers.format_excl_sheet(
        formatted, "Cross Checks", f"{tmp_path}/", "PRONET_Output.xlsx")
    workbook = load_workbook(tmp_path / "PRONET_Output.xlsx")
    try:
        sheet = workbook["Cross Checks"]
        check_id_cell = next(
            cell for cell in sheet[1] if cell.value == "Check ID")
        assert sheet.column_dimensions[check_id_cell.column_letter].hidden
    finally:
        workbook.close()
