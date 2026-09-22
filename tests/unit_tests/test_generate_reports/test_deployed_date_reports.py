"""Date findings remain visible through real V2 workbook generation/upload."""
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from generate_reports.create_trackers import CreateTrackers
from test_deployed_reporting_policy import _donor_trackers, _generator


def _date_row(network, subject, days, **overrides):
    row = {
        'network': network, 'subject': subject, 'cohort': 'CHR',
        'reports': 'Date Report', 'displayed_timepoint': 'month1',
        'displayed_form': 'visit_form', 'affected_forms': 'visit_form',
        'displayed_variable': 'visit_date', 'affected_variables': 'visit_date',
        'error_message': (
            f'visit_date : Visit date at month1 (2025-01-01) is {days} '
            'day(s) before visit_date (2025-01-11) at the earlier '
            'timepoint baseline.'),
        'currently_resolved': False, 'manually_resolved': '',
        'dates_resolved': '', 'var_translations': '',
        'time_since_last_detection': 0, 'priority_item': False,
        'comments': 'network note', 'site_comments': 'site note',
    }
    row.update(overrides)
    return row


def _real_generator(tmp_path):
    tracker, _ = _generator(tmp_path)
    # Use actual conversion and Excel serialization, with only external inputs stubbed.
    del tracker.format_excl_sheet
    del tracker.convert_to_shared_format
    style = _donor_trackers()
    tracker.colors, tracker.thin_border = style.colors, style.thin_border
    tracker.utils.can_be_float = style.utils.can_be_float
    tracker.dropbox_path = '/trackers/'
    return tracker


def test_date_report_survives_combined_site_ra_workbooks_and_upload(tmp_path):
    tracker = _real_generator(tmp_path)
    tracker.combined_tracker = pd.DataFrame([
        _date_row('PRONET', 'KC001', 10),
        _date_row('PRONET', 'KC002', 40),
        _date_row('PRESCIENT', 'PI001', 20),
        _date_row('PRESCIENT', 'ME001', 30),
        _date_row('PRONET', 'KC003', 90, reports='Old Date Report'),
        _date_row('PRONET', 'KC004', 90, affected_variables='visit_date | other_date'),
        _date_row('PRONET', 'KC005', 90, displayed_form='past_pharmaceutical_treatment'),
        _date_row('PRONET', 'KC006', 90, reports='Cross Checks'),
        _date_row('PRESCIENT', 'PI002', 90, reports='Proposed Checks'),
    ])
    tracker.collect_new_reports()
    tracker.generate_reports()
    tracker.blank_stale_sheets()

    expected = {
        'PRONET/combined/PRONET_Output_V2.xlsx': ['KC002', 'KC001'],
        'PRONET/Kansas City/PRONET_KC_Output_V2.xlsx': ['KC002', 'KC001'],
        'PRESCIENT/combined/PRESCIENT_Output_V2.xlsx': ['ME001', 'PI001'],
        'PRESCIENT/Pittsburgh/PRESCIENT_PI_Output_V2.xlsx': ['PI001'],
        'PRESCIENT/Melbourne/PRESCIENT_ME_Output_V2.xlsx': ['ME001'],
        'PRESCIENT/Melbourne/RA_One/PRESCIENT_Melbourne_Output_V2.xlsx': ['ME001'],
    }
    uploads = []
    class FakeDropbox:
        def files_upload(self, content, path, mode):
            uploads.append((content, path))
    for relative, subjects in expected.items():
        path = tmp_path / relative
        tracker.save_to_dropbox(str(path), relative, dbx=FakeDropbox())
        workbook = load_workbook(BytesIO(uploads[-1][0]))
        try:
            assert not tracker.EXCLUDED_REPORT_NAMES.intersection(workbook.sheetnames)
            rows = list(workbook['Date Report'].values)
            headers = rows[0]
            assert [row[headers.index('Participant')] for row in rows[1:]] == subjects
            assert 'Days Apart' in headers
            assert all(row[headers.index('Network Comments')] == 'network note'
                       for row in rows[1:])
            assert 'Main Report' not in workbook.sheetnames
        finally:
            workbook.close()
    assert len(uploads) == len(expected)


def test_clean_run_keeps_empty_date_tabs_and_replaces_stale_rows(tmp_path):
    tracker = _real_generator(tmp_path)
    tracker.combined_tracker = pd.DataFrame([
        _date_row('PRONET', 'KC001', 10),
        _date_row('PRESCIENT', 'ME001', 20),
    ])
    tracker.collect_new_reports()
    tracker.generate_reports()
    tracker.combined_tracker = tracker.combined_tracker.iloc[:0].copy()
    tracker.generate_reports()
    paths = list(tmp_path.rglob('*_Output_V2.xlsx'))
    assert len(paths) == 6
    for path in paths:
        workbook = load_workbook(path)
        try:
            sheet = workbook['Date Report']
            assert sheet.max_row == 1
            assert 'Participant' in next(sheet.values)
            assert 'Days Apart' in next(sheet.values)
        finally:
            workbook.close()


def test_date_report_is_created_when_a_fresh_run_has_no_date_findings(tmp_path):
    tracker = _real_generator(tmp_path)
    tracker.combined_tracker = pd.DataFrame([
        _date_row('PRONET', 'KC001', 10),
    ]).iloc[:0]
    tracker.collect_new_reports()
    tracker.generate_reports()
    assert len(list(tmp_path.rglob('*_Output_V2.xlsx'))) == 6
