"""The migration keeps deployed report surfaces and full V2 uploads.

All Dropbox interactions are in-memory fakes. Date Report is a dedicated
workbook tab; Cross/Proposed/Medication Flags remain excluded.
"""
from io import BytesIO
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.styles import Border, PatternFill, Side

from generate_reports.create_trackers import CreateTrackers
from generate_reports.calculate_resolved_errors import CalculateResolvedErrors


def _mapping():
    return {
        'subject': 'Participant', 'cohort': 'Cohort',
        'displayed_timepoint': 'Timepoint', 'displayed_form': 'Form',
        'flag_count': 'Flag Count', 'error_message': 'Flags',
        'var_translations': 'Translations',
        'time_since_last_detection': 'Days Since Detected',
        'date_resolved': 'Date Resolved', 'manually_resolved': 'Manually Resolved',
        'comments': 'Network Comments', 'site_comments': 'Site Comments',
        'priority_item': 'Priority Item',
    }


def _donor_trackers():
    tracker = object.__new__(CreateTrackers)
    tracker.colors = {color: PatternFill(fill_type='solid', fgColor='FFFFFF')
                      for color in ('green', 'yellow', 'blue', 'orange', 'red', 'grey', 'pink')}
    tracker.utils = SimpleNamespace(can_be_float=lambda value: str(value).replace('.', '', 1).isdigit())
    tracker.thin_border = Border(left=Side(style='thin'))
    return tracker


def _generator(tmp_path):
    tracker = object.__new__(CreateTrackers)
    tracker.utils = SimpleNamespace(
        pipeline_networks=('PRONET', 'PRESCIENT'),
        site_full_name_translations={'KC': 'Kansas City', 'PI': 'Pittsburgh', 'ME': 'Melbourne'})
    tracker.all_sites = {'PRONET': ['KC'], 'PRESCIENT': ['PI', 'ME']}
    tracker.all_reports = ['Main Report', 'Secondary Report']
    tracker.site_reports = ['Main Report']
    tracker.all_report_df = {}
    tracker.dropbox_output_path = str(tmp_path) + '/'
    tracker.melbourne_ras = {'RA One': ['ME001']}
    tracker.formatted_column_names = {
        network: {'combined': _mapping(), 'sites': _mapping()}
        for network in tracker.utils.pipeline_networks}
    writes = []
    tracker.format_excl_sheet = lambda frame, report, folder, filename: writes.append(
        (report, folder, filename, frame.copy()))
    tracker.convert_to_shared_format = lambda frame, network: frame.rename(
        columns={'subject': 'Participant'})
    return tracker, writes


def test_both_networks_keep_v2_reports_and_do_not_mirror_new_categories(tmp_path):
    tracker, writes = _generator(tmp_path)
    rows = []
    for network, subject in [('PRONET', 'KC001'), ('PRESCIENT', 'PI001')]:
        for report in ['Main Report', 'Secondary Report', 'Blood Report',
                       'Date Report', 'Cross Checks', 'Proposed Checks', 'Medication Flags']:
            rows.append({'network': network, 'subject': subject,
                         'reports': report, 'error_message': report})
        rows.append({'network': network, 'subject': subject,
                     'reports': 'Main Report Extended', 'error_message': 'not-main'})
    rows.append({'network': 'PRESCIENT', 'subject': 'ME001',
                 'reports': 'Non Team Forms', 'error_message': 'melbourne'})
    tracker.combined_tracker = pd.DataFrame(rows)
    original = tracker.combined_tracker.copy(deep=True)
    tracker.collect_new_reports()
    tracker.generate_reports()
    pd.testing.assert_frame_equal(tracker.combined_tracker, original)
    assert not tracker.EXCLUDED_REPORT_NAMES.intersection(report for report, *_ in writes)
    for network in ('PRONET', 'PRESCIENT'):
        combined = [entry for entry in writes if entry[2] == f'{network}_Output_V2.xlsx']
        assert {'Main Report', 'Secondary Report', 'Blood Report'} <= {entry[0] for entry in combined}
        main = next(entry[3] for entry in combined if entry[0] == 'Main Report')
        assert main['error_message'].tolist() == ['Main Report']
    assert {report for report, _, filename, _ in writes if filename == 'PRONET_KC_Output_V2.xlsx'} == {'Main Report', 'Date Report'}
    assert {report for report, _, filename, _ in writes if filename == 'PRESCIENT_PI_Output_V2.xlsx'} == {'Main Report', 'Date Report'}
    assert {report for report, _, filename, _ in writes if filename == 'PRESCIENT_ME_Output_V2.xlsx'} == {'Non Team Forms', 'Date Report'}
    assert {report for report, _, filename, _ in writes if filename == 'PRESCIENT_Melbourne_Output_V2.xlsx'} == {'Non Team Forms', 'Date Report'}


def test_explicit_new_report_names_cannot_bypass_deployed_policy(tmp_path):
    tracker, writes = _generator(tmp_path)
    tracker.all_reports = list(tracker.EXCLUDED_REPORT_NAMES)
    tracker.combined_tracker = pd.DataFrame([
        {'network': network, 'subject': 'KC001', 'reports': report}
        for network in ('PRONET', 'PRESCIENT') for report in tracker.all_reports])
    tracker.generate_reports()
    assert writes
    assert all(report == 'Date Report' and frame.empty
               for report, _, _, frame in writes)


def test_generation_respects_selected_network(tmp_path):
    tracker, writes = _generator(tmp_path)
    tracker.utils.pipeline_networks = ('PRESCIENT',)
    tracker.combined_tracker = pd.DataFrame([
        {'network': 'PRONET', 'subject': 'KC001', 'reports': 'Blood Report'},
        {'network': 'PRESCIENT', 'subject': 'PI001', 'reports': 'Main Report'}])
    tracker.collect_new_reports()
    tracker.generate_reports()
    assert 'Blood Report' not in tracker.all_reports
    assert all(filename.startswith('PRESCIENT_') for _, _, filename, _ in writes)


def test_reconciliation_reads_deployed_v2_paths_and_report_authority():
    tracker = object.__new__(CalculateResolvedErrors)
    downloads, reads, writes = [], [], []
    class FakeDropbox:
        def files_download(self, path):
            downloads.append(path)
            return None, SimpleNamespace(content=b'workbook')
    tracker.utils = SimpleNamespace(
        pipeline_networks=('PRONET', 'PRESCIENT'),
        collect_dropbox_credentials=lambda: FakeDropbox(),
        all_sites={'PRONET': ['KC'], 'PRESCIENT': ['PI', 'ME']},
        site_full_name_translations={'KC': 'Kansas City', 'PI': 'Pittsburgh', 'ME': 'Melbourne'})
    tracker.dropbox_path = '/trackers/'
    tracker.formatted_column_names = {network: {'combined': _mapping(), 'sites': _mapping()}
                                      for network in tracker.utils.pipeline_networks}
    tracker.melbourne_ras = {'RA One': ['ME001']}
    tracker._read_current = lambda: pd.DataFrame({'subject': ['KC001']})
    tracker._write_current = lambda frame: writes.append(frame)
    tracker.check_dbx_file_exists = lambda dbx, path: True
    def record(frame, columns, fields, path, dbx, network, reports, **kwargs):
        reads.append((path, fields, reports, kwargs))
        return frame
    tracker.read_dropbox_data = record
    tracker.loop_dropbox_files()
    assert downloads == ['/trackers/PRONET/combined/PRONET_Output_V2.xlsx',
                         '/trackers/PRESCIENT/combined/PRESCIENT_Output_V2.xlsx']
    assert len(writes) == 1
    assert all(path.endswith('_Output_V2.xlsx') for path, *_ in reads)
    assert not CreateTrackers.EXCLUDED_REPORT_NAMES.intersection(
        report for _, _, reports, _ in reads for report in reports)
    assert any(path.endswith('/PRONET_KC_Output_V2.xlsx') and reports == ['Main Report', 'Date Report']
               for path, _, reports, _ in reads)
    assert any(path.endswith('/RA_One/PRESCIENT_Melbourne_Output_V2.xlsx')
               and reports == ['Non Team Forms', 'Date Report'] for path, _, reports, _ in reads)
    generic = [options for _, _, _, options in reads if options.get('excl_report') is False]
    assert all(set(options['protected_reports']) == {'Main Report', 'Non Team Forms', 'Conversion Report', 'Date Report'}
               for options in generic)


@pytest.mark.parametrize('network', ['PRONET', 'PRESCIENT'])
def test_upload_replaces_the_entire_generated_workbook_without_reading_remote(tmp_path, network):
    tracker = object.__new__(CreateTrackers)
    tracker.dropbox_path = '/trackers/'
    path = tmp_path / f'{network}_Output_V2.xlsx'
    with pd.ExcelWriter(path, engine='openpyxl') as writer:
        pd.DataFrame({'Participant': ['KC001']}).to_excel(writer, sheet_name='Main Report', index=False)
        pd.DataFrame({'Participant': ['KC002']}).to_excel(writer, sheet_name='Secondary Report', index=False)
    calls = []
    class FakeDropbox:
        def files_upload(self, content, destination, mode):
            calls.append((content, destination, mode))
    tracker.save_to_dropbox(str(path), f'{network}/combined/{path.name}', dbx=FakeDropbox())
    assert len(calls) == 1
    payload, destination, mode = calls[0]
    assert payload == path.read_bytes()
    assert destination == f'/trackers/{network}/combined/{path.name}'
    assert mode.is_overwrite()
    workbook = load_workbook(BytesIO(payload))
    assert workbook.sheetnames == ['Main Report', 'Secondary Report']
    workbook.close()


def test_uploads_only_selected_network_v2_files_and_fails_on_upload_error(tmp_path):
    tracker = object.__new__(CreateTrackers)
    tracker.output_path = str(tmp_path)
    tracker.utils = SimpleNamespace(pipeline_networks=('PRONET',), collect_dropbox_credentials=lambda: object())
    paths = ['PRONET/combined/PRONET_Output_V2.xlsx',
             'PRONET/Kansas City/PRONET_KC_Output_V2.xlsx',
             'PRONET/combined/PRONET_Output.xlsx',
             'PRESCIENT/combined/PRESCIENT_Output_V2.xlsx',
             'PRONET/combined/PRONET_Output_V2.xlsx.123.tmp.xlsx']
    for relative in paths:
        target = tmp_path / 'formatted_outputs/dropbox_files' / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(b'local-workbook')
    uploads = []
    def upload(full, relative, dbx=None):
        uploads.append(relative)
        if relative.endswith('PRONET_Output_V2.xlsx'):
            raise RuntimeError('synthetic upload error')
    tracker.save_to_dropbox = upload
    with pytest.raises(RuntimeError, match='1 tracker'):
        tracker.upload_trackers()
    assert set(uploads) == set(paths[:2])


@pytest.mark.parametrize('network', ['PRONET', 'PRESCIENT'])
def test_stale_v2_tabs_are_cleared_for_both_networks_but_foreign_tabs_remain(tmp_path, network):
    tracker = _donor_trackers()
    tracker.formatted_column_names = {network: {'combined': _mapping()}}
    folder = str(tmp_path) + '/'
    filename = f'{network}_Output_V2.xlsx'
    old = pd.DataFrame({'Participant': ['old']})
    with pd.ExcelWriter(tmp_path / filename, engine='openpyxl') as writer:
        for report in ['Main Report', 'Blood Report', 'Operator Notes']:
            old.to_excel(writer, sheet_name=report, index=False)
    tracker.format_excl_sheet(pd.DataFrame({'Participant': ['new']}), 'Main Report', folder, filename)
    tracker.blank_stale_sheets()
    workbook = load_workbook(tmp_path / filename)
    assert workbook['Blood Report'].max_row == 1
    assert workbook['Operator Notes']['A2'].value == 'old'
    workbook.close()


def test_stale_sweep_does_not_touch_non_v2_workbooks(tmp_path):
    tracker = _donor_trackers()
    tracker.formatted_column_names = {'PRONET': {'combined': _mapping()}}
    path = tmp_path / 'PRONET_Output.xlsx'
    with pd.ExcelWriter(path, engine='openpyxl') as writer:
        pd.DataFrame({'Participant': ['old']}).to_excel(writer, sheet_name='Blood Report', index=False)
    before = path.read_bytes()
    tracker._sheets_written_by_path = {str(path): set()}
    tracker.blank_stale_sheets()
    assert path.read_bytes() == before


def test_format_excl_sheet_preserves_existing_sheets_on_disk(tmp_path):
    """Writing one report into an existing workbook keeps every other sheet.

    This is the local-file half of the tab-wipe contract: format_excl_sheet
    must append/replace exactly its own sheet (mode='a',
    if_sheet_exists='replace'), never recreate the file.
    """
    folder = str(tmp_path / "formatted_outputs" / "dropbox_files"
                 / "PRESCIENT" / "combined") + "/"
    filename = "PRESCIENT_Output_V2.xlsx"
    trackers = _donor_trackers()

    first = pd.DataFrame(
        [{"Participant": "ME0001", "Flags": "first-report-row"}])
    second = pd.DataFrame(
        [{"Participant": "ME0002", "Flags": "second-report-row"}])
    trackers.format_excl_sheet(first, "Main Report", folder, filename)
    trackers.format_excl_sheet(second, "Secondary Report", folder, filename)

    workbook = load_workbook(folder + filename)
    try:
        assert workbook.sheetnames == ["Main Report", "Secondary Report"]
    finally:
        workbook.close()

    # Rewriting one sheet must leave the other byte-meaningfully intact.
    updated = pd.DataFrame(
        [{"Participant": "ME0003", "Flags": "updated-main-row"}])
    trackers.format_excl_sheet(updated, "Main Report", folder, filename)

    workbook = load_workbook(folder + filename)
    try:
        assert workbook.sheetnames == ["Main Report", "Secondary Report"]
        assert list(workbook["Main Report"].values) == [
            ("Participant", "Flags"),
            ("ME0003", "updated-main-row"),
        ]
        assert list(workbook["Secondary Report"].values) == [
            ("Participant", "Flags"),
            ("ME0002", "second-report-row"),
        ]
    finally:
        workbook.close()


def test_format_excl_sheet_records_written_sheets_per_path(tmp_path):
    """The stale-sheet sweep bookkeeping tracks exactly what was written."""
    folder = str(tmp_path / "formatted_outputs" / "dropbox_files"
                 / "PRESCIENT" / "combined") + "/"
    filename = "PRESCIENT_Output_V2.xlsx"
    trackers = _donor_trackers()

    frame = pd.DataFrame(
        [{"Participant": "ME0001", "Flags": "row"}])
    trackers.format_excl_sheet(frame, "Main Report", folder, filename)
    trackers.format_excl_sheet(frame, "Scid Report", folder, filename)

    assert trackers._sheets_written_by_path == {
        folder + filename: {"Main Report", "Scid Report"},
    }



@pytest.mark.parametrize('network', ['PRONET', 'PRESCIENT'])
def test_upload_removes_excluded_tabs_from_an_existing_v2_workbook(tmp_path, network):
    tracker = object.__new__(CreateTrackers)
    tracker.output_path = str(tmp_path)
    tracker.dropbox_path = '/trackers/'
    uploaded = []
    class FakeDropbox:
        def files_upload(self, content, destination, mode):
            uploaded.append((content, destination))
    tracker.utils = SimpleNamespace(
        pipeline_networks=(network,), collect_dropbox_credentials=lambda: FakeDropbox())
    path = (tmp_path / 'formatted_outputs/dropbox_files' / network
            / 'combined' / f'{network}_Output_V2.xlsx')
    path.parent.mkdir(parents=True)
    with pd.ExcelWriter(path, engine='openpyxl') as writer:
        for report in ['Main Report', 'Secondary Report', 'Operator Notes',
                       'Date Report', 'Cross Checks', 'Proposed Checks', 'Medication Flags']:
            pd.DataFrame({'Value': [report]}).to_excel(writer, sheet_name=report, index=False)
    # No report was rewritten this run: the final upload boundary must still
    # enforce the preserved workbook contract for preexisting local files.
    tracker.upload_trackers()
    assert len(uploaded) == 1
    workbook = load_workbook(BytesIO(uploaded[0][0]))
    assert workbook.sheetnames == ['Main Report', 'Secondary Report', 'Operator Notes', 'Date Report']
    assert workbook['Date Report']['A2'].value == 'Date Report'
    assert workbook['Main Report']['A2'].value == 'Main Report'
    assert workbook['Operator Notes']['A2'].value == 'Operator Notes'
    workbook.close()


@pytest.mark.parametrize('network', ['PRONET', 'PRESCIENT'])
def test_workbook_with_only_excluded_tabs_is_not_uploaded(tmp_path, network):
    tracker = object.__new__(CreateTrackers)
    tracker.dropbox_path = '/trackers/'
    path = tmp_path / f'{network}_Output_V2.xlsx'
    with pd.ExcelWriter(path, engine='openpyxl') as writer:
        for report in tracker.EXCLUDED_REPORT_NAMES:
            pd.DataFrame({'Value': ['old']}).to_excel(writer, sheet_name=report, index=False)
    # The stub has no upload method, making any attempted publication fail.
    tracker.save_to_dropbox(str(path), f'{network}/combined/{path.name}', dbx=object())
